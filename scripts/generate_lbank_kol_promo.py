#!/usr/bin/env python3
"""Generate local LBank KOL promo output files from a table and activity copy."""

from __future__ import annotations

import argparse
import csv
import json
import posixpath
import re
import shutil
import sys
from io import BytesIO
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable
from urllib.parse import parse_qsl, quote, urlencode, urlsplit, urlunsplit
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Font


DEFAULT_OUTPUT_ROOT = Path("output")
CANONICAL_HEADERS = ["代理UID", "代理邀请码", "语言", "地区", "使用域名", "BD确认状态", "本活动是否已发送", "TG群名", "TG链接"]
OUTPUT_FIELDS = CANONICAL_HEADERS + ["注册链接", "活动链接", "生成宣发文案", "生成状态", "跳过原因", "备注"]
TEMPLATE_SHEETS = ("KOL状态 最新状态", "本期活动内容", "宣发及回链内容")
KOL_TEMPLATE_SHEET = "KOL状态 最新状态"
LEGACY_ACTIVITY_SHEET = "本期活动内容"
OUTPUT_TEMPLATE_SHEET = "宣发及回链内容"
ACTIVITY_SHEET_RE = re.compile(r"^活动(\d+)$")
TEMPLATE_OUTPUT_FIELDS = [
    "代理UID",
    "代理邀请码",
    "语言",
    "地区",
    "使用域名",
    "所属BD",
    "注册链接",
    "活动链接",
    "活动宣发内容",
    "活动宣发素材图",
    "生成状态",
    "跳过原因",
    "备注",
    "生成时间",
    "回链",
    "截图",
    "付款金额",
    "付款时间",
    "付款人",
]
INTERNAL_ROW_NUMBER_FIELD = "__source_row_number"
INTERNAL_ACTIVITY_IMAGES_FIELD = "__activity_images"
OUTPUT_IMAGE_MAX_WIDTH = 220
OUTPUT_IMAGE_MAX_HEIGHT = 120

CONFIRMED_VALUES = {"1", "true", "yes", "y", "ok", "confirmed", "已确认", "确认", "是", "bd确认"}
SENT_VALUES = {"1", "true", "yes", "y", "sent", "已发送", "已推送", "已发", "发送", "是"}

HEADER_ALIASES = {
    "代理uid": "代理UID",
    "affiliateuid": "代理UID",
    "uid": "代理UID",
    "代理邀请码": "代理邀请码",
    "邀请码": "代理邀请码",
    "invitecode": "代理邀请码",
    "refcode": "代理邀请码",
    "语言": "语言",
    "language": "语言",
    "地区": "地区",
    "region": "地区",
    "使用域名": "使用域名",
    "域名": "使用域名",
    "domain": "使用域名",
    "bd确认状态": "BD确认状态",
    "bd确认": "BD确认状态",
    "bdconfirmed": "BD确认状态",
    "投放复核状态": "投放复核状态",
    "所属bd": "所属BD",
    "本活动是否已发送": "本活动是否已发送",
    "本活动是否已经发送": "本活动是否已发送",
    "历史是否已发送": "本活动是否已发送",
    "历史是否已经发送": "本活动是否已发送",
    "是否发送": "本活动是否已发送",
    "发送状态": "本活动是否已发送",
    "sent": "本活动是否已发送",
    "tg群名": "TG群名",
    "telegram群名": "TG群名",
    "tggroup": "TG群名",
    "tg链接": "TG链接",
    "telegram链接": "TG链接",
    "tglink": "TG链接",
}

URL_RE = re.compile(r"https?://[^\s<>\"]+")
CODE_PLACEHOLDERS = ("xxxx", "XXXX", "{invite_code}", "{邀请码}", "{{invite_code}}", "{{邀请码}}")
REGISTRATION_LINK_PLACEHOLDERS = (
    "{registration_link}",
    "{ref_link}",
    "{注册链接}",
    "{邀请注册链接}",
    "{{registration_link}}",
    "{{ref_link}}",
    "{{注册链接}}",
    "{{邀请注册链接}}",
)
SECTION_MARKER_RE = re.compile(r"^\s*(?:\[([^\]]{1,64})\]|语言\s*[:：]\s*(.{1,64}))\s*$")


@dataclass(frozen=True)
class PersonalizedText:
    message: str
    registration_link: str
    activity_links: list[str]


@dataclass(frozen=True)
class TemplateActivitySource:
    sheet_name: str
    records: list[dict[str, str]]
    images_by_row: dict[int, list[Any]]
    language_pack: bool


@dataclass(frozen=True)
class CellImageAsset:
    data: bytes
    extension: str
    width: int = OUTPUT_IMAGE_MAX_WIDTH
    height: int = OUTPUT_IMAGE_MAX_HEIGHT


def normalize_header(value: object) -> str:
    raw = str(value or "").replace("\ufeff", "").strip()
    compact = re.sub(r"[\s\u3000:_-]+", "", raw).lower()
    return HEADER_ALIASES.get(compact, raw)


def normalize_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_invite_code(invite_code: str) -> str:
    code = normalize_cell(invite_code)
    code = re.sub(r"^https?://(?:www\.)?lbank\.com/ref/", "", code, flags=re.IGNORECASE)
    code = code.strip("/")
    code = re.sub(r"[\s\u3000]+", "", code)
    if not code:
        raise ValueError("缺少代理邀请码")
    return code


def normalize_domain(domain: str) -> str:
    value = normalize_cell(domain) or "lbank.com"
    if not re.match(r"^https?://", value, flags=re.IGNORECASE):
        value = "https://" + value
    return value.rstrip("/")


def build_registration_link(domain: str, invite_code: str) -> str:
    return f"{normalize_domain(domain)}/ref/{quote(normalize_invite_code(invite_code), safe='')}"


def is_confirmed(value: str) -> bool:
    return normalize_cell(value).lower() in CONFIRMED_VALUES


def is_sent(value: str) -> bool:
    return normalize_cell(value).lower() in SENT_VALUES


def strip_url_trailing_punctuation(url: str) -> tuple[str, str]:
    suffix = ""
    while url and url[-1] in "，。,.!?)）】]":
        suffix = url[-1] + suffix
        url = url[:-1]
    return url, suffix


def extract_template_codes(text: str) -> set[str]:
    codes = set()
    for match in re.finditer(r"/ref/([^?\s/#]+)", text, flags=re.IGNORECASE):
        codes.add(match.group(1).strip("/"))
    for match in re.finditer(r"[?&]icode=([^&#\s]+)", text, flags=re.IGNORECASE):
        codes.add(match.group(1))
    return {code for code in codes if code and code.lower() not in {"xxxx", "invite_code"}}


def rewrite_lbank_url(url: str, invite_code: str, registration_link: str) -> str:
    clean_url, suffix = strip_url_trailing_punctuation(url)
    parts = urlsplit(clean_url)
    if "/ref/" in parts.path:
        return registration_link + suffix
    if "/event-new/" in parts.path:
        query_pairs = [(key, value) for key, value in parse_qsl(parts.query, keep_blank_values=True) if key != "icode"]
        query_pairs.append(("icode", invite_code))
        return urlunsplit((parts.scheme, parts.netloc, parts.path, urlencode(query_pairs), parts.fragment)) + suffix
    return clean_url + suffix


def extract_activity_links(text: str) -> list[str]:
    links = []
    for match in URL_RE.finditer(text):
        url, _suffix = strip_url_trailing_punctuation(match.group(0))
        if "/event-new/" in url and url not in links:
            links.append(url)
    return links


def personalize_activity_text(text: str, invite_code: str, domain: str) -> PersonalizedText:
    code = normalize_invite_code(invite_code)
    registration_link = build_registration_link(domain, code)
    template_codes = extract_template_codes(text)

    def replace_url(match: re.Match[str]) -> str:
        return rewrite_lbank_url(match.group(0), code, registration_link)

    message = URL_RE.sub(replace_url, text)
    for placeholder in REGISTRATION_LINK_PLACEHOLDERS:
        message = message.replace(placeholder, registration_link)
    for placeholder in CODE_PLACEHOLDERS:
        message = message.replace(placeholder, code)
    for template_code in sorted(template_codes, key=len, reverse=True):
        message = message.replace(template_code, code)

    return PersonalizedText(
        message=message.strip(),
        registration_link=registration_link,
        activity_links=extract_activity_links(message),
    )


def next_run_dir(output_root: Path = DEFAULT_OUTPUT_ROOT) -> Path:
    if not output_root.exists():
        return output_root / "test1"
    max_index = 0
    for child in output_root.iterdir():
        if not child.is_dir():
            continue
        match = re.fullmatch(r"test(\d+)", child.name)
        if match:
            max_index = max(max_index, int(match.group(1)))
    return output_root / f"test{max_index + 1}"


def make_debug_dir(run_dir: Path) -> Path:
    debug_dir = run_dir / "debug"
    debug_dir.mkdir(parents=True, exist_ok=True)
    return debug_dir


def read_xlsx(path: Path) -> list[dict[str, str]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [normalize_header(value) for value in rows[0]]
        records = []
        for values in rows[1:]:
            record = {header: normalize_cell(values[index] if index < len(values) else "") for index, header in enumerate(headers) if header}
            records.append(canonicalize_record(record))
        return records
    finally:
        wb.close()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        records = []
        for row in reader:
            normalized = {normalize_header(key): normalize_cell(value) for key, value in row.items()}
            records.append(canonicalize_record(normalized))
        return records


def canonicalize_record(record: dict[str, str]) -> dict[str, str]:
    return {header: normalize_cell(record.get(header, "")) for header in CANONICAL_HEADERS}


def read_table(path: Path) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    if suffix == ".csv":
        return read_csv(path)
    raise ValueError(f"Unsupported table format: {path.suffix}. Use .xlsx or .csv")


def is_template_workbook(path: Path) -> bool:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return False
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        has_legacy_template = all(sheet_name in wb.sheetnames for sheet_name in TEMPLATE_SHEETS)
        has_activity_sheets = KOL_TEMPLATE_SHEET in wb.sheetnames and bool(activity_sheet_names(wb))
        return has_legacy_template or has_activity_sheets
    finally:
        wb.close()


def activity_sheet_names(wb: Workbook) -> list[str]:
    matches = []
    for sheet_name in wb.sheetnames:
        match = ACTIVITY_SHEET_RE.fullmatch(sheet_name)
        if match:
            matches.append((int(match.group(1)), sheet_name))
    return [sheet_name for _index, sheet_name in sorted(matches)]


def read_sheet_records(wb: Workbook, sheet_name: str) -> list[dict[str, str]]:
    ws = wb[sheet_name]
    headers = [normalize_header(cell.value) for cell in ws[1]]
    records = []
    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in values):
            continue
        record = {headers[index]: normalize_cell(values[index] if index < len(values) else "") for index in range(len(headers)) if headers[index]}
        record[INTERNAL_ROW_NUMBER_FIELD] = str(row_number)
        records.append(record)
    return records


def review_skip_reason(record: dict[str, str]) -> str | None:
    if not normalize_cell(record.get("代理邀请码", "")):
        return "缺少代理邀请码"
    if normalize_cell(record.get("投放复核状态", "")) != "通过":
        return "投放复核未通过"
    return None


def link_with_icode(link: str, invite_code: str) -> str:
    text = normalize_cell(link)
    if not text:
        return ""
    return rewrite_lbank_url(text, normalize_invite_code(invite_code), "")


def compose_template_activity_copy(activity: dict[str, str], language: str) -> str:
    title = normalize_cell(activity.get("标题", ""))
    body = normalize_cell(activity.get("正文", ""))
    event_link = normalize_cell(activity.get("链接", ""))
    parts = [part for part in [title, body] if part]
    if event_link and event_link not in body:
        parts.append(event_link)
    return "\n\n".join(parts).strip()


def template_activity_map(records: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    mapping = {}
    for record in records:
        key = normalize_language_key(record.get("语言", ""))
        if key:
            mapping[key] = record
        elif "default" not in mapping:
            mapping["default"] = record
    return mapping


def image_anchor_start(image: Any) -> tuple[int, int] | None:
    marker = getattr(getattr(image, "anchor", None), "_from", None)
    if marker is None:
        return None
    return marker.row + 1, marker.col + 1


def image_bytes(image: Any) -> bytes:
    data = getattr(image, "data", None)
    if isinstance(data, bytes):
        return data
    ref = getattr(image, "ref", None)
    if isinstance(ref, (str, Path)):
        return Path(ref).read_bytes()
    elif getattr(ref, "fp", None):
        position = ref.fp.tell()
        ref.fp.seek(0)
        data = ref.fp.read()
        ref.fp.seek(position)
        return data
    elif hasattr(ref, "seek") and hasattr(ref, "read"):
        position = ref.tell()
        ref.seek(0)
        data = ref.read()
        ref.seek(position)
        return data
    buffer = BytesIO()
    ref.save(buffer, format=getattr(image, "format", None) or "PNG")
    return buffer.getvalue()


def image_extension(image: Any) -> str:
    extension = normalize_cell(getattr(image, "extension", "")).lower()
    if extension in {".png", ".jpg", ".jpeg", ".gif", ".bmp"}:
        return extension
    image_path = normalize_cell(getattr(image, "path", ""))
    suffix = Path(image_path).suffix.lower()
    if suffix in {".png", ".jpg", ".jpeg", ".gif", ".bmp"}:
        return suffix
    image_format = normalize_cell(getattr(image, "format", "")).lower()
    if image_format in {"jpeg", "jpg"}:
        return ".jpg"
    if image_format in {"png", "gif", "bmp"}:
        return f".{image_format}"
    return ".png"


def clone_image(image: Any) -> OpenpyxlImage:
    clone = OpenpyxlImage(BytesIO(image_bytes(image)))
    clone.width = image.width
    clone.height = image.height
    scale_image_to_fit(clone, OUTPUT_IMAGE_MAX_WIDTH, OUTPUT_IMAGE_MAX_HEIGHT)
    return clone


def scale_image_to_fit(image: OpenpyxlImage, max_width: int, max_height: int) -> None:
    if image.width <= 0 or image.height <= 0:
        return
    scale = min(max_width / image.width, max_height / image.height, 1)
    image.width = int(image.width * scale)
    image.height = int(image.height * scale)


def xlsx_part_join(base_part: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def xlsx_rels_part(part_name: str) -> str:
    directory = posixpath.dirname(part_name)
    filename = posixpath.basename(part_name)
    return posixpath.join(directory, "_rels", f"{filename}.rels")


def parse_relationship_targets(xml_data: bytes) -> dict[str, str]:
    rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    root = ET.fromstring(xml_data)
    return {
        relationship.attrib.get("Id", ""): relationship.attrib.get("Target", "")
        for relationship in root.findall(f"{rel_ns}Relationship")
        if relationship.attrib.get("Id")
    }


def workbook_sheet_parts(workbook_path: Path) -> dict[str, str]:
    main_ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    try:
        with ZipFile(workbook_path) as archive:
            workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
            relationships = parse_relationship_targets(archive.read("xl/_rels/workbook.xml.rels"))
    except (KeyError, OSError, ET.ParseError):
        return {}

    parts = {}
    for sheet in workbook_root.findall("main:sheets/main:sheet", main_ns):
        sheet_name = sheet.attrib.get("name", "")
        relationship_id = sheet.attrib.get(rel_attr, "")
        target = relationships.get(relationship_id, "")
        if sheet_name and target:
            parts[sheet_name] = xlsx_part_join("xl/workbook.xml", target)
    return parts


def read_wps_cell_image_assets(workbook_path: Path) -> dict[str, CellImageAsset]:
    drawing_ns = {
        "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }
    embed_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    try:
        with ZipFile(workbook_path) as archive:
            cellimages_part = "xl/cellimages.xml"
            relationships = parse_relationship_targets(archive.read(xlsx_rels_part(cellimages_part)))
            root = ET.fromstring(archive.read(cellimages_part))
            assets = {}
            for pic in root.findall(".//xdr:pic", drawing_ns):
                properties = pic.find(".//xdr:cNvPr", drawing_ns)
                blip = pic.find(".//a:blip", drawing_ns)
                if properties is None or blip is None:
                    continue
                image_id = properties.attrib.get("name", "")
                relationship_id = blip.attrib.get(embed_attr, "")
                target = relationships.get(relationship_id, "")
                if not image_id or not target or target.upper() == "NULL":
                    continue
                media_part = xlsx_part_join(cellimages_part, target)
                extension = Path(media_part).suffix.lower() or ".png"
                assets[image_id] = CellImageAsset(
                    data=archive.read(media_part),
                    extension=extension,
                    width=OUTPUT_IMAGE_MAX_WIDTH,
                    height=OUTPUT_IMAGE_MAX_HEIGHT,
                )
            return assets
    except (KeyError, OSError, ET.ParseError):
        return {}


def worksheet_dispimg_images_by_row(workbook_path: Path, sheet_name: str, poster_column: int | None) -> dict[int, list[CellImageAsset]]:
    sheet_part = workbook_sheet_parts(workbook_path).get(sheet_name)
    if not sheet_part:
        return {}
    cell_images = read_wps_cell_image_assets(workbook_path)
    if not cell_images:
        return {}

    main_ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    images_by_row: dict[int, list[CellImageAsset]] = {}
    try:
        with ZipFile(workbook_path) as archive:
            root = ET.fromstring(archive.read(sheet_part))
    except (KeyError, OSError, ET.ParseError):
        return {}

    for cell in root.findall(".//main:c", main_ns):
        coordinate = cell.attrib.get("r", "")
        match = re.fullmatch(r"([A-Z]+)(\d+)", coordinate)
        if not match:
            continue
        column_letters, row_text = match.groups()
        column_number = 0
        for letter in column_letters:
            column_number = column_number * 26 + ord(letter) - ord("A") + 1
        if poster_column and column_number != poster_column:
            continue

        formula = cell.find("main:f", main_ns)
        value = cell.find("main:v", main_ns)
        formula_text = "\n".join(
            part
            for part in [
                formula.text if formula is not None else "",
                value.text if value is not None else "",
            ]
            if part
        )
        seen_image_ids = set()
        for image_id in re.findall(r"DISPIMG\(\s*[\"']([^\"']+)[\"']", formula_text, flags=re.IGNORECASE):
            if image_id in seen_image_ids:
                continue
            seen_image_ids.add(image_id)
            image = cell_images.get(image_id)
            if image:
                images_by_row.setdefault(int(row_text), []).append(image)
    return images_by_row


def merge_images_by_row(target: dict[int, list[Any]], source: dict[int, list[Any]]) -> None:
    for row_number, images in source.items():
        target.setdefault(row_number, []).extend(images)


def template_activity_image_map(wb: Workbook, records: list[dict[str, str]], sheet_name: str = LEGACY_ACTIVITY_SHEET) -> dict[str, list[Any]]:
    ws = wb[sheet_name]
    headers = [normalize_header(cell.value) for cell in ws[1]]
    poster_column = headers.index("配图") + 1 if "配图" in headers else None
    images_by_row: dict[int, list[Any]] = {}

    for image in getattr(ws, "_images", []):
        start = image_anchor_start(image)
        if not start:
            continue
        row_number, column_number = start
        if poster_column and column_number != poster_column:
            continue
        images_by_row.setdefault(row_number, []).append(image)

    mapping: dict[str, list[Any]] = {}
    for record in records:
        row_number = int(record.get(INTERNAL_ROW_NUMBER_FIELD, "0") or 0)
        images = images_by_row.get(row_number, [])
        if not images:
            continue
        key = normalize_language_key(record.get("语言", ""))
        if key:
            mapping[key] = images
        elif "default" not in mapping:
            mapping["default"] = images
    return mapping


def template_activity_images_by_row(wb: Workbook, sheet_name: str = LEGACY_ACTIVITY_SHEET, workbook_path: Path | None = None) -> dict[int, list[Any]]:
    ws = wb[sheet_name]
    headers = [normalize_header(cell.value) for cell in ws[1]]
    poster_column = headers.index("配图") + 1 if "配图" in headers else None
    images_by_row: dict[int, list[Any]] = {}

    for image in getattr(ws, "_images", []):
        start = image_anchor_start(image)
        if not start:
            continue
        row_number, column_number = start
        if poster_column and column_number != poster_column:
            continue
        images_by_row.setdefault(row_number, []).append(image)
    if workbook_path:
        merge_images_by_row(images_by_row, worksheet_dispimg_images_by_row(workbook_path, sheet_name, poster_column))
    return images_by_row


def replace_activity_sheet_image_cells_with_asset_paths(wb: Workbook, sheet_name: str, asset_paths_by_row: dict[int, list[str]]) -> None:
    if sheet_name not in wb.sheetnames:
        return
    ws = wb[sheet_name]
    headers = [normalize_header(cell.value) for cell in ws[1]]
    if "配图" not in headers:
        return
    poster_column = headers.index("配图") + 1
    for row_number, asset_paths in asset_paths_by_row.items():
        if row_number <= 1 or row_number > ws.max_row or not asset_paths:
            continue
        ws.cell(row=row_number, column=poster_column).value = " | ".join(asset_paths)


def save_activity_image_assets(activity_index: int, images: list[Any], assets_dir: Path, asset_key: str = "") -> list[str]:
    if not images:
        return []
    assets_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    clean_key = re.sub(r"[^A-Za-z0-9_-]+", "_", normalize_cell(asset_key)).strip("_")
    name_key = f"_{clean_key}" if clean_key else ""
    for image_index, image in enumerate(images, start=1):
        target = assets_dir / f"activity_{activity_index}{name_key}_image_{image_index}{image_extension(image)}"
        target.write_bytes(image_bytes(image))
        paths.append(f"assets/{target.name}")
    return paths


def load_template_activity_sources(wb: Workbook, workbook_path: Path | None = None) -> list[TemplateActivitySource]:
    activity_sheets = activity_sheet_names(wb)
    if activity_sheets:
        return [
            TemplateActivitySource(
                sheet_name=sheet_name,
                records=read_sheet_records(wb, sheet_name),
                images_by_row=template_activity_images_by_row(wb, sheet_name, workbook_path),
                language_pack=True,
            )
            for sheet_name in activity_sheets
        ]

    records = read_sheet_records(wb, LEGACY_ACTIVITY_SHEET) if LEGACY_ACTIVITY_SHEET in wb.sheetnames else []
    images_by_row = template_activity_images_by_row(wb, LEGACY_ACTIVITY_SHEET, workbook_path) if LEGACY_ACTIVITY_SHEET in wb.sheetnames else {}
    return [
        TemplateActivitySource(
            sheet_name=LEGACY_ACTIVITY_SHEET,
            records=[record],
            images_by_row=images_by_row,
            language_pack=False,
        )
        for record in (records or [{}])
    ]


def select_activity_record(
    source: TemplateActivitySource,
    kol_language: str,
) -> tuple[dict[str, str], str | None, str | None]:
    if not source.language_pack:
        return source.records[0], None, None

    mapping = template_activity_map(source.records)
    language_key = normalize_language_key(kol_language)
    if language_key in mapping:
        return mapping[language_key], None, None
    if "en" in mapping:
        return mapping["en"], "EN", None
    if "default" in mapping:
        return mapping["default"], "default", None
    return {}, None, "缺少对应语言活动文案"


def remove_template_output_sheets(wb: Workbook) -> None:
    for sheet_name in list(wb.sheetnames):
        if sheet_name == "宣发及回链内容" or sheet_name.startswith("宣发及回链内容_活动"):
            del wb[sheet_name]


def template_output_sheet_name(activity_index: int, activity_count: int) -> str:
    if activity_count == 1:
        return "宣发及回链内容"
    return f"宣发及回链内容_活动{activity_index}"


def ensure_template_output_sheet(ws) -> None:
    ws.delete_rows(1, ws.max_row)
    ws.append(TEMPLATE_OUTPUT_FIELDS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    width_by_field = {
        "代理UID": 18,
        "代理邀请码": 18,
        "语言": 12,
        "地区": 12,
        "使用域名": 22,
        "所属BD": 18,
        "注册链接": 34,
        "活动链接": 58,
        "活动宣发内容": 90,
        "活动宣发素材图": 32,
        "生成状态": 14,
        "跳过原因": 24,
        "备注": 34,
        "生成时间": 22,
        "回链": 42,
        "截图": 32,
        "付款金额": 14,
        "付款时间": 20,
        "付款人": 16,
    }
    for index, field in enumerate(TEMPLATE_OUTPUT_FIELDS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width_by_field.get(field, 18)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_template_output_sheet(wb: Workbook, rows: list[dict[str, Any]], sheet_name: str = "宣发及回链内容") -> None:
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
    ensure_template_output_sheet(ws)
    image_column = TEMPLATE_OUTPUT_FIELDS.index("活动宣发素材图") + 1
    for output_row_number, row in enumerate(rows, start=2):
        ws.append([row.get(field, "") for field in TEMPLATE_OUTPUT_FIELDS])
        for image in row.get(INTERNAL_ACTIVITY_IMAGES_FIELD, []):
            cloned = clone_image(image)
            cloned.anchor = ws.cell(row=output_row_number, column=image_column).coordinate
            ws.add_image(cloned)
            ws.row_dimensions[output_row_number].height = max(
                ws.row_dimensions[output_row_number].height or 0,
                cloned.height * 0.75,
            )
    for body_row in ws.iter_rows(min_row=2):
        for cell in body_row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.auto_filter.ref = ws.dimensions


def run_template_generation(input_table: Path, output_root: Path = DEFAULT_OUTPUT_ROOT) -> dict[str, object]:
    wb = load_workbook(input_table)
    kol_records = read_sheet_records(wb, KOL_TEMPLATE_SHEET)
    activity_sources = load_template_activity_sources(wb, input_table)

    output_root.mkdir(parents=True, exist_ok=True)
    run_dir = next_run_dir(output_root)
    run_dir.mkdir(parents=True, exist_ok=False)
    debug_dir = make_debug_dir(run_dir)
    assets_dir = run_dir / "assets"
    copy_input_table(input_table, debug_dir)

    generated_rows: list[dict[str, str]] = []
    skipped_rows: list[dict[str, str]] = []
    activity_asset_paths: list[str] = []
    fallback_rows: list[dict[str, str]] = []
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    remove_template_output_sheets(wb)
    activity_count = len(activity_sources)
    for activity_index, source in enumerate(activity_sources, start=1):
        sheet_rows: list[dict[str, Any]] = []
        saved_assets_by_row: dict[int, list[str]] = {}
        for activity in source.records:
            activity_row_number = int(activity.get(INTERNAL_ROW_NUMBER_FIELD, "0") or 0)
            if activity_row_number in saved_assets_by_row:
                continue
            asset_key = f"row_{activity_row_number}" if source.language_pack else ""
            saved_assets_by_row[activity_row_number] = save_activity_image_assets(
                activity_index,
                source.images_by_row.get(activity_row_number, []),
                assets_dir,
                asset_key=asset_key,
            )
            activity_asset_paths.extend(saved_assets_by_row[activity_row_number])
        replace_activity_sheet_image_cells_with_asset_paths(wb, source.sheet_name, saved_assets_by_row)

        for record in kol_records:
            base = {
                "代理UID": normalize_cell(record.get("代理UID", "")),
                "代理邀请码": normalize_cell(record.get("代理邀请码", "")),
                "语言": normalize_cell(record.get("语言", "")),
                "地区": normalize_cell(record.get("地区", "")),
                "使用域名": normalize_cell(record.get("使用域名", "")),
                "所属BD": normalize_cell(record.get("所属BD", "")),
                "回链": "",
                "截图": "",
                "付款金额": "",
                "付款时间": "",
                "付款人": "",
            }
            reason = review_skip_reason(record)
            activity, fallback_language, missing_activity_reason = select_activity_record(source, record.get("语言", ""))
            if not reason and missing_activity_reason:
                reason = missing_activity_reason
            if not reason and not any(normalize_cell(activity.get(field, "")) for field in ("标题", "正文", "链接")):
                reason = "缺少活动文案"
            activity_row_number = int(activity.get(INTERNAL_ROW_NUMBER_FIELD, "0") or 0)
            asset_paths = saved_assets_by_row.get(activity_row_number, [])

            remark = f"语言 {record.get('语言', '')} 未提供，已使用 {fallback_language} 英文兜底" if fallback_language else ""

            if reason:
                row = {**base, "注册链接": "", "活动链接": "", "活动宣发内容": "", "活动宣发素材图": "", "生成状态": "已跳过", "跳过原因": reason, "备注": remark, "生成时间": generated_at}
                skipped_rows.append({**row, "BD确认状态": record.get("投放复核状态", ""), "本活动是否已发送": "", "TG群名": "", "TG链接": "", "生成宣发文案": "", "备注": row["备注"]})
                sheet_rows.append(row)
                continue

            text = compose_template_activity_copy(activity, record.get("语言", ""))
            personalized = personalize_activity_text(text, record["代理邀请码"], record.get("使用域名", ""))
            event_link = link_with_icode(activity.get("链接", ""), record["代理邀请码"])
            activity_links = event_link or " | ".join(personalized.activity_links)
            asset_text = " | ".join(asset_paths) if asset_paths else normalize_cell(activity.get("配图", ""))
            row = {
                **base,
                "注册链接": personalized.registration_link,
                "活动链接": activity_links,
                "活动宣发内容": personalized.message,
                "活动宣发素材图": asset_text,
                "生成状态": "已生成",
                "跳过原因": "",
                "备注": remark,
                "生成时间": generated_at,
            }
            generated_rows.append(
                {
                    "代理UID": row["代理UID"],
                    "代理邀请码": row["代理邀请码"],
                    "语言": row["语言"],
                    "地区": row["地区"],
                    "使用域名": row["使用域名"],
                    "BD确认状态": record.get("投放复核状态", ""),
                    "本活动是否已发送": "",
                    "TG群名": "",
                    "TG链接": "",
                    "注册链接": row["注册链接"],
                    "活动链接": row["活动链接"],
                    "生成宣发文案": row["活动宣发内容"],
                    "生成状态": "已生成",
                    "跳过原因": "",
                    "备注": row["备注"],
                }
            )
            if fallback_language:
                fallback_rows.append(
                    {
                        "activity": source.sheet_name,
                        "代理UID": row["代理UID"],
                        "代理邀请码": row["代理邀请码"],
                        "原语言": row["语言"],
                        "兜底语言": fallback_language,
                    }
                )
            sheet_rows.append(row)

        write_template_output_sheet(wb, sheet_rows, template_output_sheet_name(activity_index, activity_count))

    filled_template = run_dir / "filled_template.xlsx"
    wb.save(filled_template)
    wb.close()

    write_generated_markdown(debug_dir / "generated_messages.md", generated_rows)
    write_csv_rows(debug_dir / "generated_messages.csv", OUTPUT_FIELDS, generated_rows + skipped_rows)
    write_xlsx_rows(debug_dir / "generated_messages.xlsx", OUTPUT_FIELDS, generated_rows + skipped_rows)
    write_csv_rows(debug_dir / "skipped_rows.csv", OUTPUT_FIELDS, skipped_rows)
    write_xlsx_rows(debug_dir / "skipped_rows.xlsx", OUTPUT_FIELDS, skipped_rows, sheet_name="跳过记录")

    summary = {
        "run_dir": str(run_dir),
        "input_table": str(input_table),
        "mode": "activity_sheet_template" if any(source.language_pack for source in activity_sources) else "three_sheet_template",
        "activity_text": ", ".join(source.sheet_name for source in activity_sources),
        "activity_count": activity_count,
        "generated_count": len(generated_rows),
        "skipped_count": len(skipped_rows),
        "fallback_count": len(fallback_rows),
        "fallback_rows": fallback_rows,
        "output_files": [
            "filled_template.xlsx",
            *activity_asset_paths,
            "debug/input_original.xlsx",
            "debug/generated_messages.xlsx",
            "debug/generated_messages.csv",
            "debug/generated_messages.md",
            "debug/skipped_rows.xlsx",
            "debug/skipped_rows.csv",
        ],
    }
    return summary


def split_activity_templates(text: str) -> dict[str, str]:
    templates: dict[str, list[str]] = {}
    current_key = "default"
    current_lines: list[str] = []
    saw_marker = False

    def flush() -> None:
        content = "\n".join(current_lines).strip()
        if content:
            templates[current_key] = current_lines.copy()

    for line in text.splitlines():
        match = SECTION_MARKER_RE.match(line)
        if match:
            flush()
            saw_marker = True
            current_key = normalize_language_key(match.group(1) or match.group(2) or "")
            current_lines = []
            continue
        current_lines.append(line)
    flush()

    if not saw_marker:
        return {"default": text}
    return {key: "\n".join(lines).strip() for key, lines in templates.items()}


def make_activity_loader(
    activity_path: Path | None = None,
    activity_copy: str | None = None,
) -> tuple[Callable[[str], str | None], Callable[[Path], None], str]:
    if activity_copy is not None:
        templates = split_activity_templates(activity_copy)

        def load_from_copy(language: str) -> str | None:
            return templates.get(normalize_language_key(language)) or templates.get("default")

        def backup_copy(run_dir: Path) -> None:
            (run_dir / "activity_text.txt").write_text(activity_copy, encoding="utf-8")

        return load_from_copy, backup_copy, "inline_activity_copy"

    if activity_path is None:
        raise ValueError("activity_text or activity_copy is required")

    if activity_path.is_file():
        text = activity_path.read_text(encoding="utf-8")
        templates = split_activity_templates(text)

        def load_single_or_sections(language: str) -> str | None:
            return templates.get(normalize_language_key(language)) or templates.get("default")

        def backup_single(run_dir: Path) -> None:
            shutil.copy2(activity_path, run_dir / "activity_text.txt")

        return load_single_or_sections, backup_single, str(activity_path)

    if activity_path.is_dir():
        templates = {normalize_language_key(path.stem): path.read_text(encoding="utf-8") for path in activity_path.glob("*.txt")}

        def load_by_language(language: str) -> str | None:
            return templates.get(normalize_language_key(language)) or templates.get("default")

        def backup_dir(run_dir: Path) -> None:
            shutil.copytree(activity_path, run_dir / "activity_texts")

        return load_by_language, backup_dir, str(activity_path)

    raise FileNotFoundError(f"Activity text path not found: {activity_path}")


def normalize_language_key(language: str) -> str:
    value = normalize_cell(language).lower().replace("_", "-")
    return re.sub(r"[\s\u3000]+", "", value)


def skip_reason(record: dict[str, str]) -> str | None:
    if not normalize_cell(record.get("代理邀请码", "")):
        return "缺少代理邀请码"
    if not is_confirmed(record.get("BD确认状态", "")):
        return "BD未确认"
    if is_sent(record.get("本活动是否已发送", "")):
        return "已发送"
    return None


def write_generated_markdown(path: Path, generated_rows: list[dict[str, str]]) -> None:
    parts = []
    for row in generated_rows:
        title = f"## 代理UID: {row['代理UID']} | TG群名: {row['TG群名']}".rstrip()
        meta = f"TG链接: {row['TG链接']}" if row["TG链接"] else ""
        parts.append("\n".join(part for part in [title, meta, "", row["生成宣发文案"]] if part != ""))
    path.write_text("\n\n".join(parts).strip() + "\n", encoding="utf-8")


def write_csv_rows(path: Path, fieldnames: list[str], rows: list[dict[str, str]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx_rows(path: Path, fieldnames: list[str], rows: list[dict[str, str]], sheet_name: str = "生成结果") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(fieldnames)
    for row in rows:
        ws.append([row.get(field, "") for field in fieldnames])

    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = header_alignment
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = body_alignment

    width_by_field = {
        "代理UID": 14,
        "代理邀请码": 16,
        "语言": 12,
        "地区": 12,
        "使用域名": 22,
        "BD确认状态": 14,
        "本活动是否已发送": 18,
        "TG群名": 22,
        "TG链接": 32,
        "注册链接": 34,
        "活动链接": 58,
        "生成宣发文案": 80,
        "生成状态": 12,
        "跳过原因": 18,
        "备注": 28,
    }
    for index, field in enumerate(fieldnames, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width_by_field.get(field, 18)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def copy_input_table(input_table: Path, run_dir: Path) -> None:
    suffix = input_table.suffix.lower()
    target = run_dir / ("input_original.xlsx" if suffix in {".xlsx", ".xlsm"} else "input_original.csv")
    shutil.copy2(input_table, target)


def run_generation(
    input_table: Path,
    activity_text: Path | None = None,
    output_root: Path = DEFAULT_OUTPUT_ROOT,
    activity_copy: str | None = None,
) -> dict[str, object]:
    if activity_text is None and activity_copy is None and is_template_workbook(input_table):
        return run_template_generation(input_table, output_root)

    records = read_table(input_table)
    load_activity_text, backup_activity, activity_source = make_activity_loader(activity_text, activity_copy)

    output_root.mkdir(parents=True, exist_ok=True)
    run_dir = next_run_dir(output_root)
    run_dir.mkdir(parents=True, exist_ok=False)
    debug_dir = make_debug_dir(run_dir)

    copy_input_table(input_table, debug_dir)
    backup_activity(debug_dir)

    generated_rows: list[dict[str, str]] = []
    skipped_rows: list[dict[str, str]] = []
    output_rows: list[dict[str, str]] = []

    for record in records:
        reason = skip_reason(record)
        if reason:
            row = {**record, "注册链接": "", "活动链接": "", "生成宣发文案": "", "生成状态": "已跳过", "跳过原因": reason, "备注": ""}
            skipped_rows.append(row)
            output_rows.append(row)
            continue

        text = load_activity_text(record.get("语言", ""))
        if not text:
            row = {**record, "注册链接": "", "活动链接": "", "生成宣发文案": "", "生成状态": "已跳过", "跳过原因": "缺少对应语言活动文案", "备注": ""}
            skipped_rows.append(row)
            output_rows.append(row)
            continue

        try:
            personalized = personalize_activity_text(text, record["代理邀请码"], record["使用域名"])
        except ValueError as exc:
            row = {**record, "注册链接": "", "活动链接": "", "生成宣发文案": "", "生成状态": "已跳过", "跳过原因": str(exc), "备注": ""}
            skipped_rows.append(row)
            output_rows.append(row)
            continue

        row = {
            **record,
            "注册链接": personalized.registration_link,
            "活动链接": " | ".join(personalized.activity_links),
            "生成宣发文案": personalized.message,
            "生成状态": "已生成",
            "跳过原因": "",
            "备注": "",
        }
        generated_rows.append(row)
        output_rows.append(row)

    write_generated_markdown(debug_dir / "generated_messages.md", generated_rows)
    write_csv_rows(debug_dir / "generated_messages.csv", OUTPUT_FIELDS, output_rows)
    write_xlsx_rows(run_dir / "generated_messages.xlsx", OUTPUT_FIELDS, output_rows)
    write_csv_rows(debug_dir / "skipped_rows.csv", OUTPUT_FIELDS, skipped_rows)
    write_xlsx_rows(debug_dir / "skipped_rows.xlsx", OUTPUT_FIELDS, skipped_rows, sheet_name="跳过记录")

    summary = {
        "run_dir": str(run_dir),
        "input_table": str(input_table),
        "activity_text": activity_source,
        "generated_count": len(generated_rows),
        "skipped_count": len(skipped_rows),
        "output_files": [
            "generated_messages.xlsx",
            "debug/input_original.xlsx",
            "debug/activity_text.txt" if activity_text is None or activity_text.is_file() else "debug/activity_texts",
            "debug/generated_messages.md",
            "debug/generated_messages.csv",
            "debug/skipped_rows.xlsx",
            "debug/skipped_rows.csv",
        ],
    }
    return summary


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate local LBank KOL promo outputs.")
    parser.add_argument("--input-table", required=True, type=Path, help="Input .xlsx or .csv table.")
    parser.add_argument("--activity-text", type=Path, help="Activity text .txt file, or folder of language .txt files.")
    parser.add_argument("--activity-copy", help="Activity copy pasted directly in the command.")
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT, help="Root output folder. Defaults to ./output.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    if not args.activity_text and args.activity_copy is None and not is_template_workbook(args.input_table):
        print("One of --activity-text or --activity-copy is required.", file=sys.stderr)
        return 2
    summary = run_generation(args.input_table, args.activity_text, output_root=args.output_root, activity_copy=args.activity_copy)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
