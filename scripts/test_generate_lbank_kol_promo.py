import base64
import csv
import json
import sys
import tempfile
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from PIL import Image as PILImage

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from generate_lbank_kol_promo import DEFAULT_OUTPUT_ROOT, build_registration_link, is_template_workbook, next_run_dir, personalize_activity_text, run_generation


HEADERS = [
    "代理UID",
    "代理邀请码",
    "语言",
    "地区",
    "使用域名",
    "BD确认状态",
    "本活动是否已发送",
    "TG群名",
    "TG链接",
    "生成宣发文案",
    "生成状态",
    "跳过原因",
]


def write_workbook(path: Path, rows: list[dict[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "KOL"
    ws.append(HEADERS)
    for row in rows:
        ws.append([row.get(header, "") for header in HEADERS])
    wb.save(path)


TINY_PNG = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8"
    "/x8AAwMB/aznF3QAAAAASUVORK5CYII="
)


def write_tiny_png(path: Path) -> None:
    path.write_bytes(base64.b64decode(TINY_PNG))


def write_test_png(path: Path, size: tuple[int, int] = (1200, 400)) -> None:
    image = PILImage.new("RGB", size, "purple")
    image.save(path)


def add_wps_cell_image(path: Path, image_path: Path, image_id: str = "ID_TEST_IMAGE") -> None:
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    sheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", rel_ns)
    ET.register_namespace("", sheet_ns)
    with ZipFile(path, "r") as source:
        entries = {info.filename: source.read(info.filename) for info in source.infolist()}

    workbook_rels_part = "xl/_rels/workbook.xml.rels"
    rels_root = ET.fromstring(entries[workbook_rels_part])
    existing_ids = [
        int(relationship.attrib["Id"][3:])
        for relationship in rels_root
        if relationship.attrib.get("Id", "").startswith("rId") and relationship.attrib["Id"][3:].isdigit()
    ]
    next_id = max(existing_ids or [0]) + 1
    rels_root.append(
        ET.Element(
            f"{{{rel_ns}}}Relationship",
            {
                "Id": f"rId{next_id}",
                "Type": "http://www.wps.cn/officeDocument/2020/cellImage",
                "Target": "cellimages.xml",
            },
        )
    )
    entries[workbook_rels_part] = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)
    entries["xl/cellimages.xml"] = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<etc:cellImages xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:etc="http://www.wps.cn/officeDocument/2017/etCustomData">
  <etc:cellImage>
    <xdr:pic>
      <xdr:nvPicPr>
        <xdr:cNvPr id="1" name="{image_id}"/>
        <xdr:cNvPicPr><a:picLocks noChangeAspect="1"/></xdr:cNvPicPr>
      </xdr:nvPicPr>
      <xdr:blipFill>
        <a:blip r:embed="rId1"/>
        <a:stretch><a:fillRect/></a:stretch>
      </xdr:blipFill>
      <xdr:spPr>
        <a:xfrm><a:off x="0" y="0"/><a:ext cx="11430000" cy="3810000"/></a:xfrm>
        <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
      </xdr:spPr>
    </xdr:pic>
  </etc:cellImage>
</etc:cellImages>
""".encode("utf-8")
    entries["xl/_rels/cellimages.xml.rels"] = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>
</Relationships>
"""
    entries["xl/media/image1.png"] = image_path.read_bytes()
    for name, data in list(entries.items()):
        if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
            continue
        root = ET.fromstring(data)
        changed = False
        for cell in root.findall(f".//{{{sheet_ns}}}c"):
            formula = cell.find(f"{{{sheet_ns}}}f")
            if formula is None or not formula.text or "DISPIMG" not in formula.text:
                continue
            value = cell.find(f"{{{sheet_ns}}}v")
            if value is None:
                value = ET.SubElement(cell, f"{{{sheet_ns}}}v")
            value.text = f"={formula.text}"
            changed = True
        if changed:
            entries[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    temp_path = path.with_suffix(".tmp.xlsx")
    with ZipFile(temp_path, "w") as target:
        for name, data in entries.items():
            target.writestr(name, data)
    temp_path.replace(path)


def write_template_workbook(path: Path, embedded_image: Path | None = None, poster_value: str = "poster.png") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "KOL状态 最新状态"
    ws.append(["代理UID", "代理邀请码", "语言", "地区", "使用域名", "所属BD", "BD申请状态", "投放复核状态"])
    ws.append(["TEST_UID_001", "ABC123", "EN", "PHP", "LBank.com", "BD_TEST", "申请", "通过"])
    ws.append(["TEST_UID_002", "ZZZ999", "EN", "PHP", "LBank.com", "BD_TEST", "申请", "待确认"])

    ws = wb.create_sheet("本期活动内容")
    ws.append(["语言", "标题", "正文", "链接", "配图"])
    ws.append(
        [
            "EN",
            "Activity 5 10,000 USDT In WOJAK Rewards & Futures Bonus",
            "Join here:\nhttps://www.lbank.com/event-new/10001753-wojak-sausage",
            "https://www.lbank.com/event-new/10001753-wojak-sausage",
            poster_value,
        ]
    )
    if embedded_image:
        ws.add_image(OpenpyxlImage(str(embedded_image)), "E2")

    ws = wb.create_sheet("宣发及回链内容")
    ws.append(
        [
            "代理UID",
            "代理邀请码",
            "语言",
            "地区",
            "使用域名",
            "所属BD",
            "注册链接",
            "活动链接",
            "活动宣发内容",
            "活动宣发素材图",
            "生成状态",
            "跳过原因",
            "备注",
            "生成时间",
            "回链",
            "截图",
            "付款金额",
            "付款时间",
            "付款人",
        ]
    )
    wb.save(path)


def write_multi_activity_template_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "KOL状态 最新状态"
    ws.append(["代理UID", "代理邀请码", "语言", "地区", "使用域名", "所属BD", "BD申请状态", "投放复核状态"])
    ws.append(["TEST_UID_001", "ABC123", "VN", "PHP", "LBank.com", "BD_TEST", "申请", "通过"])
    ws.append(["TEST_UID_002", "ZZZ999", "EN", "PHP", "LBank.com", "BD_TEST", "申请", "待确认"])

    ws = wb.create_sheet("本期活动内容")
    ws.append(["标题", "正文", "链接", "配图"])
    ws.append(
        [
            "Activity One",
            "First activity body\nhttps://www.lbank.com/event-new/10000001-first",
            "https://www.lbank.com/event-new/10000001-first",
            "",
        ]
    )
    ws.append(
        [
            "Activity Two",
            "Second activity body\nhttps://www.lbank.com/event-new/10000002-second",
            "https://www.lbank.com/event-new/10000002-second",
            "",
        ]
    )

    ws = wb.create_sheet("宣发及回链内容")
    ws.append(["代理UID", "代理邀请码", "语言", "地区", "使用域名", "所属BD", "活动宣发内容", "活动宣发素材图"])
    wb.save(path)


def write_language_pack_template_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "KOL状态 最新状态"
    ws.append(["代理UID", "代理邀请码", "代理备注", "代理层级", "语言", "地区", "域名", "所属BD", "BD申请状态", "投放复核状态"])
    ws.append(["TEST_UID_EN", "EN123", "English KOL", "总代理", "EN", "PHP", "LBank.com", "BD_TEST", "申请", "通过"])
    ws.append(["TEST_UID_TR", "TR456", "Turkish KOL", "总代理", "TR", "TR", "LBank.com", "BD_TEST", "申请", "通过"])

    ws = wb.create_sheet("本期活动内容")
    ws.append(["标题", "正文", "链接", "配图"])
    ws.append(["Wrong Legacy Title", "Wrong legacy body", "https://www.lbank.com/event-new/99999999-wrong", ""])

    ws = wb.create_sheet("宣发及回链内容")
    ws.append(["代理UID", "代理邀请码", "语言", "地区", "域名", "所属BD", "活动宣发内容", "活动宣发素材图", "回链", "截图", "付款金额", "付款时间", "付款人"])

    ws = wb.create_sheet("活动1")
    ws.append(["语言", "标题", "正文", "链接", "配图"])
    ws.append(["EN", "English Activity", "English body", "https://www.lbank.com/event-new/10001812-0fee-stocks", ""])
    ws.append(["CN", "繁中活動", "繁中正文", "https://www.lbank.com/zh-TC/event-new/10001812-0fee-stocks", ""])
    wb.save(path)


class GenerateLbankKolPromoTest(unittest.TestCase):
    def test_default_output_root_uses_portable_output_folder(self):
        self.assertEqual(DEFAULT_OUTPUT_ROOT, Path("output"))

    def test_build_registration_link_uses_row_domain_and_invite_code(self):
        self.assertEqual(
            build_registration_link("backup.lbank.info", "38O7X"),
            "https://backup.lbank.info/ref/38O7X",
        )
        self.assertEqual(
            build_registration_link("https://www.lbank.com/", "AB 12"),
            "https://www.lbank.com/ref/AB12",
        )
        self.assertEqual(
            build_registration_link("https://www.lbank.com/", "AB\u300012"),
            "https://www.lbank.com/ref/AB12",
        )

    def test_personalize_activity_text_removes_spaces_inside_invite_codes(self):
        text = "Step 2: Register -> https://www.lbank.com/event-new/10001469-deposit-lossprotection-allusers"

        result = personalize_activity_text(text, invite_code="AB C 003", domain="lbank.com")

        self.assertEqual(result.registration_link, "https://lbank.com/ref/ABC003")
        self.assertIn("icode=ABC003", result.message)
        self.assertNotIn("AB+C+003", result.message)

    def test_personalize_activity_text_replaces_step1_and_step2_codes(self):
        text = """超級禮包：充值&交易獎勵🎁
👉步驟1: 邀請碼註冊鏈接 →https://lbank.com/ref/38O7X
👉步驟2: 點擊活動鏈接 → https://www.lbank.com/event-new/10001465-first-deposit-trade-newusers?icode=38O7X
"""

        result = personalize_activity_text(
            text,
            invite_code="NEW99",
            domain="lbank.com",
        )

        self.assertIn("https://lbank.com/ref/NEW99", result.message)
        self.assertIn("first-deposit-trade-newusers?icode=NEW99", result.message)
        self.assertNotIn("38O7X", result.message)
        self.assertEqual(result.registration_link, "https://lbank.com/ref/NEW99")
        self.assertEqual(
            result.activity_links,
            ["https://www.lbank.com/event-new/10001465-first-deposit-trade-newusers?icode=NEW99"],
        )

    def test_personalize_activity_text_adds_icode_when_event_link_has_none(self):
        text = "Step 2: Register the event -> https://www.lbank.com/event-new/10001469-deposit-lossprotection-allusers"

        result = personalize_activity_text(text, invite_code="CODE7", domain="lbank.com")

        self.assertIn(
            "https://www.lbank.com/event-new/10001469-deposit-lossprotection-allusers?icode=CODE7",
            result.message,
        )

    def test_next_run_dir_uses_incrementing_test_folder_names(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            (root / "test1").mkdir()
            (root / "test3").mkdir()

            self.assertEqual(next_run_dir(root).name, "test4")

    def test_run_generation_reads_xlsx_filters_rows_and_writes_outputs(self):
        activity_text = """[繁中]
🎉超級禮包：充值&交易獎勵
👉步驟1: 邀請碼註冊鏈接 →{注册链接}
👉步驟2: 點擊活動鏈接 → https://www.lbank.com/event-new/10001469-deposit-lossprotection-allusers

[en]
🎉Triple Protection Mega Bonus Event
Step 1: Sign up with code {invite_code}
→{registration_link}
Step 2: Register the event
→ https://www.lbank.com/event-new/10001469-deposit-lossprotection-allusers
"""
        rows = [
            {
                "代理UID": "A001",
                "代理邀请码": "AAA111",
                "语言": "繁中",
                "地区": "HK",
                "使用域名": "lbank.com",
                "BD确认状态": "已确认",
                "本活动是否已发送": "",
                "TG群名": "Alpha",
                "TG链接": "https://t.me/alpha",
            },
            {
                "代理UID": "B002",
                "代理邀请码": "BBB222",
                "语言": "en",
                "地区": "PH",
                "使用域名": "www.lbank.com",
                "BD确认状态": "待确认",
                "本活动是否已发送": "",
                "TG群名": "Beta",
                "TG链接": "https://t.me/beta",
            },
            {
                "代理UID": "C003",
                "代理邀请码": "CCC333",
                "语言": "en",
                "地区": "PH",
                "使用域名": "lbank.com",
                "BD确认状态": "已确认",
                "本活动是否已发送": "已发送",
                "TG群名": "Gamma",
                "TG链接": "https://t.me/gamma",
            },
            {
                "代理UID": "D004",
                "代理邀请码": "DDD444",
                "语言": "en",
                "地区": "PH",
                "使用域名": "www.lbank.com",
                "BD确认状态": "已确认",
                "本活动是否已发送": "否",
                "TG群名": "Delta",
                "TG链接": "https://t.me/delta",
            },
        ]

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            workbook_path = tmp / "input.xlsx"
            output_root = tmp / "output"
            write_workbook(workbook_path, rows)

            summary = run_generation(workbook_path, activity_copy=activity_text, output_root=output_root)

            run_dir = output_root / "test1"
            self.assertEqual(summary["run_dir"], str(run_dir))
            self.assertEqual(summary["generated_count"], 2)
            self.assertEqual(summary["skipped_count"], 2)
            debug_dir = run_dir / "debug"
            self.assertEqual(sorted(path.name for path in run_dir.iterdir()), ["debug", "generated_messages.xlsx"])
            self.assertTrue((run_dir / "generated_messages.xlsx").exists())
            self.assertTrue((debug_dir / "generated_messages.md").exists())
            self.assertTrue((debug_dir / "generated_messages.csv").exists())
            self.assertTrue((debug_dir / "skipped_rows.xlsx").exists())
            self.assertTrue((debug_dir / "skipped_rows.csv").exists())
            self.assertFalse((debug_dir / "run_summary.json").exists())
            self.assertTrue((debug_dir / "input_original.xlsx").exists())
            self.assertTrue((debug_dir / "activity_text.txt").exists())

            self.assertEqual(
                summary["output_files"],
                [
                    "generated_messages.xlsx",
                    "debug/input_original.xlsx",
                    "debug/activity_text.txt",
                    "debug/generated_messages.md",
                    "debug/generated_messages.csv",
                    "debug/skipped_rows.xlsx",
                    "debug/skipped_rows.csv",
                ],
            )

            md = (debug_dir / "generated_messages.md").read_text(encoding="utf-8")
            self.assertIn("代理UID: A001", md)
            self.assertIn("https://lbank.com/ref/AAA111", md)
            self.assertIn("deposit-lossprotection-allusers?icode=AAA111", md)
            self.assertIn("Sign up with code DDD444", md)
            self.assertIn("https://www.lbank.com/ref/DDD444", md)
            self.assertNotIn("{注册链接}", md)
            self.assertNotIn("{invite_code}", md)

            with (debug_dir / "generated_messages.csv").open(encoding="utf-8-sig", newline="") as handle:
                output_rows = list(csv.DictReader(handle))
            self.assertEqual([row["代理UID"] for row in output_rows], ["A001", "B002", "C003", "D004"])
            self.assertEqual(output_rows[0]["注册链接"], "https://lbank.com/ref/AAA111")
            self.assertEqual(output_rows[0]["生成状态"], "已生成")
            self.assertEqual(output_rows[1]["生成状态"], "已跳过")

            skipped = (debug_dir / "skipped_rows.csv").read_text(encoding="utf-8-sig")
            self.assertIn("BD未确认", skipped)
            self.assertIn("已发送", skipped)

            output_wb = load_workbook(run_dir / "generated_messages.xlsx", data_only=True)
            ws = output_wb.active
            headers = [cell.value for cell in ws[1]]
            self.assertIn("本活动是否已发送", headers)
            self.assertIn("生成宣发文案", headers)
            self.assertIn("生成状态", headers)
            self.assertIn("跳过原因", headers)
            self.assertIn("备注", headers)
            self.assertEqual(ws.max_row, 5)
            output_wb.close()

    def test_run_generation_fills_three_sheet_template_workbook(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            workbook_path = tmp / "template.xlsx"
            output_root = tmp / "output"
            write_template_workbook(workbook_path)

            summary = run_generation(workbook_path, output_root=output_root)

            run_dir = output_root / "test1"
            output_path = run_dir / "filled_template.xlsx"
            self.assertEqual(summary["generated_count"], 1)
            self.assertEqual(summary["skipped_count"], 1)
            self.assertTrue(output_path.exists())
            debug_dir = run_dir / "debug"
            self.assertEqual(sorted(path.name for path in run_dir.iterdir()), ["debug", "filled_template.xlsx"])
            self.assertTrue((debug_dir / "input_original.xlsx").exists())
            self.assertTrue((debug_dir / "generated_messages.xlsx").exists())
            self.assertTrue((debug_dir / "generated_messages.csv").exists())
            self.assertTrue((debug_dir / "generated_messages.md").exists())
            self.assertTrue((debug_dir / "skipped_rows.xlsx").exists())
            self.assertTrue((debug_dir / "skipped_rows.csv").exists())
            self.assertFalse((debug_dir / "run_summary.json").exists())
            self.assertEqual(
                summary["output_files"],
                [
                    "filled_template.xlsx",
                    "debug/input_original.xlsx",
                    "debug/generated_messages.xlsx",
                    "debug/generated_messages.csv",
                    "debug/generated_messages.md",
                    "debug/skipped_rows.xlsx",
                    "debug/skipped_rows.csv",
                ],
            )

            wb = load_workbook(output_path, data_only=True)
            ws = wb["宣发及回链内容"]
            headers = [cell.value for cell in ws[1]]
            row2 = {headers[index]: ws.cell(row=2, column=index + 1).value for index in range(len(headers))}
            row3 = {headers[index]: ws.cell(row=3, column=index + 1).value for index in range(len(headers))}

            self.assertEqual(row2["代理UID"], "TEST_UID_001")
            self.assertEqual(row2["注册链接"], "https://LBank.com/ref/ABC123")
            self.assertEqual(row2["活动链接"], "https://www.lbank.com/event-new/10001753-wojak-sausage?icode=ABC123")
            self.assertIn("Activity 5 10,000 USDT", row2["活动宣发内容"])
            self.assertIn("https://www.lbank.com/event-new/10001753-wojak-sausage?icode=ABC123", row2["活动宣发内容"])
            self.assertNotIn("https://LBank.com/ref/ABC123", row2["活动宣发内容"])
            self.assertNotIn("Step 1", row2["活动宣发内容"])
            self.assertNotIn("步骤1", row2["活动宣发内容"])
            self.assertNotIn("步驟1", row2["活动宣发内容"])
            self.assertEqual(row2["活动宣发素材图"], "poster.png")
            self.assertEqual(row2["生成状态"], "已生成")
            self.assertIn(row2["备注"], (None, ""))
            self.assertIsNotNone(row2["生成时间"])

            self.assertEqual(row3["代理UID"], "TEST_UID_002")
            self.assertEqual(row3["生成状态"], "已跳过")
            self.assertEqual(row3["跳过原因"], "投放复核未通过")
            self.assertIn(row3["备注"], (None, ""))
            wb.close()

    def test_run_generation_writes_one_output_sheet_per_activity(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            workbook_path = tmp / "template.xlsx"
            output_root = tmp / "output"
            write_multi_activity_template_workbook(workbook_path)

            summary = run_generation(workbook_path, output_root=output_root)

            self.assertEqual(summary["activity_count"], 2)
            self.assertEqual(summary["generated_count"], 2)
            self.assertEqual(summary["skipped_count"], 2)

            wb = load_workbook(output_root / "test1" / "filled_template.xlsx", data_only=True)
            self.assertIn("宣发及回链内容_活动1", wb.sheetnames)
            self.assertIn("宣发及回链内容_活动2", wb.sheetnames)
            self.assertNotIn("宣发及回链内容", wb.sheetnames)

            ws1 = wb["宣发及回链内容_活动1"]
            ws2 = wb["宣发及回链内容_活动2"]
            headers1 = [cell.value for cell in ws1[1]]
            headers2 = [cell.value for cell in ws2[1]]
            row1_activity1 = {headers1[index]: ws1.cell(row=2, column=index + 1).value for index in range(len(headers1))}
            row1_activity2 = {headers2[index]: ws2.cell(row=2, column=index + 1).value for index in range(len(headers2))}

            self.assertEqual(ws1.max_row, 3)
            self.assertEqual(ws2.max_row, 3)
            self.assertIn("10000001-first?icode=ABC123", row1_activity1["活动链接"])
            self.assertIn("10000002-second?icode=ABC123", row1_activity2["活动链接"])
            self.assertIn("Activity One", row1_activity1["活动宣发内容"])
            self.assertIn("Activity Two", row1_activity2["活动宣发内容"])
            self.assertEqual(ws1.cell(row=3, column=headers1.index("生成状态") + 1).value, "已跳过")
            self.assertEqual(ws2.cell(row=3, column=headers2.index("生成状态") + 1).value, "已跳过")
            wb.close()

    def test_run_generation_prefers_activity_sheet_language_pack_with_en_fallback(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            workbook_path = tmp / "template.xlsx"
            output_root = tmp / "output"
            write_language_pack_template_workbook(workbook_path)

            self.assertTrue(is_template_workbook(workbook_path))
            summary = run_generation(workbook_path, output_root=output_root)

            self.assertEqual(summary["mode"], "activity_sheet_template")
            self.assertEqual(summary["activity_text"], "活动1")
            self.assertEqual(summary["activity_count"], 1)
            self.assertEqual(summary["generated_count"], 2)
            self.assertEqual(summary["skipped_count"], 0)
            self.assertEqual(summary["fallback_count"], 1)
            self.assertEqual(summary["fallback_rows"][0]["原语言"], "TR")
            self.assertEqual(summary["fallback_rows"][0]["兜底语言"], "EN")

            wb = load_workbook(output_root / "test1" / "filled_template.xlsx", data_only=True)
            ws = wb["宣发及回链内容"]
            headers = [cell.value for cell in ws[1]]
            row_en = {headers[index]: ws.cell(row=2, column=index + 1).value for index in range(len(headers))}
            row_tr = {headers[index]: ws.cell(row=3, column=index + 1).value for index in range(len(headers))}

            self.assertIn("English Activity", row_en["活动宣发内容"])
            self.assertIn("English Activity", row_tr["活动宣发内容"])
            self.assertNotIn("Wrong Legacy Title", row_en["活动宣发内容"])
            self.assertNotIn("Wrong Legacy Title", row_tr["活动宣发内容"])
            self.assertIn("10001812-0fee-stocks?icode=TR456", row_tr["活动链接"])
            self.assertIn("10001812-0fee-stocks?icode=TR456", row_tr["活动宣发内容"])
            self.assertNotIn("https://LBank.com/ref/TR456", row_tr["活动宣发内容"])
            self.assertNotIn("Step 1", row_tr["活动宣发内容"])
            self.assertNotIn("步骤1", row_tr["活动宣发内容"])
            self.assertNotIn("步驟1", row_tr["活动宣发内容"])
            self.assertIn(row_en["备注"], (None, ""))
            self.assertEqual(row_tr["备注"], "语言 TR 未提供，已使用 EN 英文兜底")
            wb.close()

    def test_run_generation_saves_embedded_activity_image_as_asset_path(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            workbook_path = tmp / "template.xlsx"
            output_root = tmp / "output"
            image_path = tmp / "poster.png"
            write_test_png(image_path)
            write_template_workbook(workbook_path, embedded_image=image_path, poster_value="")

            run_generation(workbook_path, output_root=output_root)

            wb = load_workbook(output_root / "test1" / "filled_template.xlsx", data_only=True)
            ws = wb["宣发及回链内容"]
            headers = [cell.value for cell in ws[1]]
            image_col = headers.index("活动宣发素材图")

            asset_path = ws.cell(row=2, column=image_col + 1).value
            self.assertEqual(len(ws._images), 0)
            self.assertEqual(asset_path, "assets/activity_1_image_1.png")
            self.assertTrue((output_root / "test1" / asset_path).exists())
            wb.close()

    def test_run_generation_saves_wps_dispimg_cell_image_as_asset_path(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            workbook_path = tmp / "template.xlsx"
            output_root = tmp / "output"
            image_path = tmp / "poster.png"
            write_test_png(image_path)
            write_language_pack_template_workbook(workbook_path)

            wb = load_workbook(workbook_path)
            ws = wb["活动1"]
            ws["E2"] = '=DISPIMG("ID_TEST_IMAGE",1)'
            wb.save(workbook_path)
            wb.close()
            add_wps_cell_image(workbook_path, image_path)

            run_generation(workbook_path, output_root=output_root)

            wb = load_workbook(output_root / "test1" / "filled_template.xlsx", data_only=True)
            ws = wb["宣发及回链内容"]
            headers = [cell.value for cell in ws[1]]
            image_col = headers.index("活动宣发素材图")

            asset_path = ws.cell(row=2, column=image_col + 1).value
            fallback_asset_path = ws.cell(row=3, column=image_col + 1).value
            self.assertEqual(asset_path, "assets/activity_1_row_2_image_1.png")
            self.assertEqual(fallback_asset_path, "assets/activity_1_row_2_image_1.png")
            self.assertTrue((output_root / "test1" / asset_path).exists())
            wb.close()

            output_path = output_root / "test1" / "filled_template.xlsx"
            source_wb = load_workbook(output_path, data_only=False)
            source_ws = source_wb["活动1"]
            self.assertIn("DISPIMG", source_ws["E2"].value)
            source_wb.close()
            with ZipFile(output_path) as archive:
                names = archive.namelist()
                self.assertIn("xl/cellimages.xml", names)
                self.assertIn("xl/_rels/cellimages.xml.rels", names)
                workbook_rels = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")
                self.assertIn("http://www.wps.cn/officeDocument/2020/cellImage", workbook_rels)


if __name__ == "__main__":
    unittest.main()
